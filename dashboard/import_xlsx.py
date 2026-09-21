"""Einmaliger Import der DATEV-Excel (August 2026) in report.db."""
import sys
import openpyxl

sys.path.insert(0, str(__import__("pathlib").Path(__file__).parent))
from report_db import connect

XLSX = "/Users/dwfb/Downloads/BWA_Scialdone_August_2026.xlsx"
MONTH = "2026-08"

VV_FIELD_MAP = {
    "Umsatzerlöse": "umsatz",
    "Bestandsveränderg. FE/UE": "bestandsveraenderung",
    "Aktivierte Eigenleistungen": "aktivierteEigenleistungen",
    "Material-/Wareneinkauf": "material",
    "So. betr. Erlöse": "sonstBetrErloese",
    "Personalkosten": "personal",
    "Raumkosten": "raum",
    "Betriebliche Steuern": "betrSteuern",
    "Versicherungen/Beiträge": "versicherungen",
    "Besondere Kosten": "besondereKosten",
    "Fahrzeugkosten (ohne Steuer)": "fahrzeug",
    "Werbe-/Reisekosten": "werbungReise",
    "Kosten Warenabgabe": "kostenWarenabgabe",
    "Abschreibungen": "abschreibungen",
    "Reparatur/Instandhaltung": "reparatur",
    "Sonstige Kosten": "sonstigeKosten",
    "Zinsaufwand": "zinsaufwand",
    "Sonstiger neutraler Aufwand": "sonstNeutralAufwand",
    "Zinserträge": "zinsertraege",
    "Sonstiger neutraler Ertrag": "sonstNeutralErtrag",
    "Steuern Einkommen u. Ertrag": "steuernEinkommenErtrag",
}


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    con = connect()
    cur = con.cursor()
    cur.execute("DELETE FROM susa_entries WHERE month=?", (MONTH,))
    cur.execute("DELETE FROM opos_entries WHERE month=?", (MONTH,))

    n = 0
    blatt_map = {
        "SuSa Blatt 1 (Kl. 0-3)": (1, 0),
        "SuSa Blatt 2 (Kl. 3-6)": (2, 0),
        "SuSa Blatt 3 (Kl. 6-9)": (3, 0),
        "SuSa Blatt 4 (Kreditoren)": (4, 1),
    }
    for sheet, (blatt, is_kred) in blatt_map.items():
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=6):
            a, b = row[0].value, row[1].value
            if a is None or b is None:
                continue  # Zwischensaldo-Zeilen ohne Konto
            konto = str(int(a)) if isinstance(a, (int, float)) else str(a).strip()
            beschriftung = str(b).strip()
            if beschriftung.lower().startswith("summe"):
                continue
            cur.execute(
                "INSERT INTO susa_entries (month, blatt, konto, beschriftung, eb_wert, "
                "eb_sh, soll_monat, haben_monat, soll_kum, haben_kum, saldo_sh, is_kreditor) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                (MONTH, blatt, konto, beschriftung,
                 num(row[2].value), str(row[3].value or "S").strip()[:1],
                 num(row[4].value), num(row[5].value),
                 num(row[6].value), num(row[7].value),
                 str(row[9].value or "S").strip()[:1], is_kred))
            n += 1

    # Vorjahr-Spalte für den Vorjahresvergleich (Aug/2025)
    cur.execute("DELETE FROM vv_prev WHERE month=?", ("2025-08",))
    ws = wb["Vorjahresvergleich"]
    v = 0
    for row in ws.iter_rows(min_row=6):
        label = row[0].value
        if isinstance(label, str) and label.strip() in VV_FIELD_MAP and row[2].value is not None:
            cur.execute("INSERT OR REPLACE INTO vv_prev (month, field, amount) VALUES (?,?,?)",
                        ("2025-08", VV_FIELD_MAP[label.strip()], num(row[2].value)))
            v += 1
    con.commit()
    print(f"susa_rows={n} vv_prev_rows={v}")


if __name__ == "__main__":
    main()
